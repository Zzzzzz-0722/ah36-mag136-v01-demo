"""Read source values without evaluating, repairing or inventing observations."""
from pathlib import Path
import hashlib
import json
import math

import openpyxl
import pandas as pd

SHEETS = {'工艺参数库': '记录ID', '焊接Case': 'CaseID', '逐道参数': 'PassID', '实验结果': 'ResultID'}
NUM = ['Thickness', 'GrooveAngle', 'RootGap']
CAT = ['Position', 'PassRole']
FEATURES = NUM + CAT
V02_MODEL_FEATURES = NUM + ['PassRole']
TARGETS = {'Current': '推荐_实测电流_A', 'Voltage': '推荐_实测电压_V', 'TravelSpeed': '推荐_实测焊速_mm_min'}
ROLES = ['Root', 'Fill', 'Cap']
PROJECT_SCOPE = {
    'Material': 'AH36',
    'Process': 'MAG 136',
    'JointType': 'Single-V Butt Joint',
    'Position': 'PA',
    'Thickness': (10.0, 20.0),
    'GrooveAngle': (40.0, 60.0),
    'RootGap': (5.0, 8.0),
}


def number(value):
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (TypeError, ValueError):
        return None


def position(value):
    return {'PA / AWS 1G': 'PA', '1G (flat)': 'PA', '平对接': 'PA',
            'Vertical up': 'PF', '3G Uphill': 'PF', '45°斜立对接（向上）': 'Inclined45Up'}.get(value)


def roles(value):
    if value == 'Root/Fill/Cap共用固定电流电压':
        return ROLES
    if value and ('打底' in value or value == 'Root'):
        return ['Root']
    if value in ['填充', 'Fill']:
        return ['Fill']
    if value and ('盖面' in value or value == 'Cap'):
        return ['Cap']
    return [None]


def standardized_process(case):
    method = str(case.get('Case_焊接方法') or '')
    iso = number(case.get('Case_ISO4063'))
    if 'SAW' in method or '埋弧' in method:
        return 'Mixed MAG/SAW'
    if iso != 136:
        return None if iso is None else f'ISO {int(iso)}'
    return 'MAG 136'


def standardized_joint(case):
    joint = str(case.get('Case_接头类型') or '')
    groove = str(case.get('Case_坡口形式') or '')
    if '不等厚' in groove:
        return 'Unequal-thickness Butt Joint'
    is_butt = any(x in joint.lower() for x in ['对接', 'butt', 'plate to plate'])
    is_single_v = any(x in groove.lower() for x in ['单v', 'single v', 'v型'])
    return 'Single-V Butt Joint' if is_butt and is_single_v else None


def in_range(value, bounds):
    return value is not None and bounds[0] <= value <= bounds[1]


def missing_to_none(value):
    return None if value is None or pd.isna(value) else value


def evidence_suggestion(row, missing_fields, incompatible_fields):
    nature = str(row.get('Pass_参数性质') or '')
    verified = str(row.get('Pass_核验状态') or '').startswith(('已核验', '已核对'))
    source_clear = bool(row.get('Case_来源名称') or row.get('Case_来源链接'))
    if incompatible_fields:
        return 'E', 'Known project identity or geometry is incompatible with V0.2 scope'
    if len(missing_fields) in [1, 2] and any(number(row.get('Pass_' + field)) not in [None, 0]
                                               for field in TARGETS.values()):
        return 'D', 'Useful target value exists but one or two required context fields are missing'
    if any(x in nature for x in ['候选', 'WPS', '工程', '设定']) and verified and source_clear:
        return 'B', 'Traceable engineering or WPS setting; human review remains open'
    if '范围' in nature or '窗口' in nature:
        return 'C', 'Range or window retained for Constraint Check only'
    if any(x in nature for x in ['实验', '实测', 'PQR']) and verified and source_clear:
        return 'A', 'Measured or experimental setting with traceable source metadata'
    return 'D', 'Potentially useful record with incomplete evidence or process identity'


def add_v02_eligibility(training, formulas):
    """Add V0.2 scope and target-specific eligibility without altering source columns."""
    source_groups = sorted(x for x in training.SourceGroup.dropna().unique())
    source_ids = {value: f'SRC-{i:03d}' for i, value in enumerate(source_groups, 1)}
    rows = []
    for original in training.to_dict('records'):
        row = dict(original)
        row['Material'] = row.get('Case_材料牌号')
        row['Process'] = standardized_process(row)
        row['JointType'] = standardized_joint(row)
        required = {
            'Material': missing_to_none(row['Material']), 'Process': missing_to_none(row['Process']),
            'JointType': missing_to_none(row['JointType']), 'Position': missing_to_none(row.get('Position')),
            'Thickness': number(row.get('Thickness')),
            'GrooveAngle': number(row.get('GrooveAngle')), 'RootGap': number(row.get('RootGap')),
            'PassRole': missing_to_none(row.get('PassRole')),
        }
        missing = [key for key, value in required.items() if value is None]
        incompatible = []
        for key in ['Material', 'Process', 'JointType', 'Position']:
            if required[key] is not None and required[key] != PROJECT_SCOPE[key]:
                incompatible.append(key)
        for key in NUM:
            if required[key] is not None and not in_range(required[key], PROJECT_SCOPE[key]):
                incompatible.append(key)
        scope_missing = [key for key in missing if key != 'PassRole']
        row['InProjectScope'] = not scope_missing and not incompatible
        row['ScopeMissingFields'] = ', '.join(scope_missing)
        row['ScopeMismatchFields'] = ', '.join(incompatible)
        level, basis = evidence_suggestion(row, missing, incompatible)
        row['EvidenceLevel_Suggested'] = level
        row['EvidenceLevel_Basis'] = basis
        row['EvidenceLevel_Manual'] = ''
        row['EvidenceReviewStatus'] = 'Pending'
        row['SourceID'] = source_ids.get(row.get('SourceGroup'), '')
        row['Source'] = row.get('Case_来源名称') or row.get('SourceGroup') or ''
        base_reasons = []
        if not row['InProjectScope']:
            if scope_missing:
                base_reasons.append('Missing project scope fields: ' + ','.join(scope_missing))
            if incompatible:
                base_reasons.append('Outside project scope: ' + ','.join(incompatible))
        if row.get('PassRole') not in ROLES:
            base_reasons.append('Missing or ambiguous PassRole')
        if level not in ['A', 'B']:
            base_reasons.append(f'EvidenceLevel {level} is not eligible for supervised training')
        if row.get('Case_训练准入') == '否' or row.get('Pass_训练准入') == '否':
            base_reasons.append('Source explicitly excludes training')
        for target, field in TARGETS.items():
            value = number(row.get('Pass_' + field))
            formula = ('逐道参数', row.get('Pass_ExcelRow'), field) in formulas
            reasons = list(base_reasons)
            if value is None or value <= 0:
                reasons.append('Missing or invalid target')
            if formula:
                reasons.append('Formula-based target excluded')
            eligible = not reasons
            row['Eligible_' + target] = eligible
            row['EligibilityReason_' + target] = '; '.join(reasons)
        rows.append(row)
    return pd.DataFrame(rows)


def build_promotion_queue(frame):
    candidates = []
    for _, row in frame.drop_duplicates('PassID').iterrows():
        missing = [x for x in (str(row.ScopeMissingFields) + ',' +
                               ('PassRole' if row.PassRole not in ROLES else '')).split(',') if x.strip()]
        missing = [x.strip() for x in missing]
        values = {target: number(row.get('Pass_' + field)) for target, field in TARGETS.items()}
        valid_targets = [target for target, value in values.items() if value is not None and value > 0]
        if row.ScopeMismatchFields or not 1 <= len(missing) <= 2 or not valid_targets:
            continue
        priority = 'High' if len(missing) == 1 and len(valid_targets) >= 2 else ('Medium' if len(valid_targets) >= 2 else 'Low')
        action = 'Verify and populate ' + ', '.join(missing) + '; retain source values unchanged'
        candidates.append({
            'Priority': priority, 'CaseID': row.CaseID, 'PassID': row.PassID,
            'Current DataClass': row.DataClass, 'EvidenceLevel': row.EvidenceLevel_Suggested,
            'Current': values['Current'], 'Voltage': values['Voltage'], 'TravelSpeed': values['TravelSpeed'],
            'MissingFields': ', '.join(missing),
            'ExclusionReason': '; '.join(dict.fromkeys(
                reason for value in [row.get('EligibilityReason_Current'), row.get('EligibilityReason_Voltage'),
                                     row.get('EligibilityReason_TravelSpeed')]
                for reason in str(value or '').split('; ') if reason)),
            'SourceID': row.SourceID, 'Source': row.Source, 'RecommendedAction': action,
        })
    queue = pd.DataFrame(candidates)
    if not queue.empty:
        queue['_order'] = queue.Priority.map({'High': 0, 'Medium': 1, 'Low': 2})
        queue = queue.sort_values(['_order', 'CaseID', 'PassID']).drop(columns='_order').reset_index(drop=True)
    return queue


def v02_audit(frame, queue, source_summary):
    passes = frame.drop_duplicates('PassID')
    reason_counts = {}
    for _, row in passes.iterrows():
        reasons = set()
        for target in TARGETS:
            if not row['Eligible_' + target]:
                reasons.update(reason for reason in str(row['EligibilityReason_' + target]).split('; ')
                               if reason and reason != 'nan')
        for reason in reasons:
            if reason:
                reason_counts[reason] = reason_counts.get(reason, 0) + 1
    return {
        'version': 'V0.2-data-architecture',
        'source_sha256': source_summary['sha256'],
        'project_scope': PROJECT_SCOPE,
        'model_features_v02': V02_MODEL_FEATURES,
        'counts': {
            'total_cases': int(source_summary['焊接Case']['records']),
            'total_passes': int(passes.PassID.nunique()),
            'in_project_scope_passes': int(passes.InProjectScope.sum()),
            'evidence_levels': passes.EvidenceLevel_Suggested.value_counts().sort_index().to_dict(),
            'eligible_current': int(passes.Eligible_Current.sum()),
            'eligible_voltage': int(passes.Eligible_Voltage.sum()),
            'eligible_travel_speed': int(passes.Eligible_TravelSpeed.sum()),
            'promotion_candidates': int(len(queue)),
        },
        'context_counts': {
            'rows_after_role_expansion': int(len(frame)),
            'eligible_current': int(frame.Eligible_Current.sum()),
            'eligible_voltage': int(frame.Eligible_Voltage.sum()),
            'eligible_travel_speed': int(frame.Eligible_TravelSpeed.sum()),
        },
        'main_exclusion_reasons': dict(sorted(reason_counts.items(), key=lambda item: (-item[1], item[0]))),
        'notes': [
            'Counts under counts use unique PassID; context_counts include explicit Root/Fill/Cap expansion.',
            'EvidenceLevel is suggested from metadata. EvidenceLevel_Manual remains blank for human review.',
            'Range midpoints are never converted to supervised targets.',
            'No model was trained or overwritten by the V0.2 audit pipeline.',
        ],
    }


def read_source(path):
    raw = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    tables, issues, formula_cells = {}, [], set()
    summary = {}
    for name, key in SHEETS.items():
        s, c = raw[name], cached[name]
        headers = [v.value for v in s[1]]
        records = []
        for row in s.iter_rows(min_row=2):
            # Ignore formatting-only rows and empty template formulas without a key.
            if not row[0].value:
                if any(x.value is not None and x.data_type != 'f' for x in row):
                    issues.append({'Sheet': name, 'Cell': f'A{row[0].row}', 'Issue': 'missing primary key'})
                continue
            record = {h: c.cell(row[0].row, i + 1).value for i, h in enumerate(headers) if h}
            record['ExcelRow'] = row[0].row
            for i, cell in enumerate(row):
                if cell.data_type == 'f':
                    formula_cells.add((name, row[0].row, headers[i]))
                    issues.append({'Sheet': name, 'Cell': cell.coordinate, 'RecordID': row[0].value,
                                   'Issue': 'formula in source', 'Formula': cell.value,
                                   'CachedValue': c[cell.coordinate].value, 'Field': headers[i]})
            records.append(record)
        df = pd.DataFrame(records, columns=[h for h in headers if h] + ['ExcelRow'])
        if df[key].duplicated().any():
            raise ValueError(f'Duplicate {key} in {name}')
        tables[name] = df
        summary[name] = {'records': len(df), 'fields': len(df.columns) - 1,
                         'missing': df.isna().sum().to_dict(),
                         'field_names': list(df.columns),
                         'formatted_rows': s.max_row}
    ids = set(tables['焊接Case']['CaseID'])
    for name in ['逐道参数', '实验结果']:
        orphans = set(tables[name]['CaseID']) - ids
        if orphans:
            raise ValueError(f'Orphan CaseID in {name}: {orphans}')
    passes = tables['逐道参数'].set_index('PassID')
    for r in tables['实验结果'].to_dict('records'):
        pid = r.get('PassID_可空')
        if pd.notna(pid) and (pid not in passes.index or passes.loc[pid, 'CaseID'] != r['CaseID']):
            raise ValueError(f'Invalid result PassID: {pid}')
    # Validate all bounds and cached single values, including constraints.
    for name, df in tables.items():
        for r in df.to_dict('records'):
            for lo in [k for k in r if '下限' in k]:
                hi = lo.replace('下限', '上限')
                a, b = number(r.get(lo)), number(r.get(hi))
                if a is not None and b is not None and a > b:
                    issues.append({'Sheet': name, 'RecordID': r[SHEETS[name]], 'Issue': f'reversed bounds: {lo}'})
            if name == '逐道参数':
                for label, value, lo, hi in [('Current', '推荐_实测电流_A', '电流下限_A', '电流上限_A'),
                                            ('Voltage', '推荐_实测电压_V', '电压下限_V', '电压上限_V'),
                                            ('TravelSpeed', '推荐_实测焊速_mm_min', '焊速下限_mm_min', '焊速上限_mm_min')]:
                    v, a, b = number(r.get(value)), number(r.get(lo)), number(r.get(hi))
                    if v is not None and ((a is not None and v < a) or (b is not None and v > b)):
                        issues.append({'Sheet': name, 'RecordID': r['PassID'], 'Issue': f'{label} outside own bounds', 'CachedValue': v})
    summary['sha256'] = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    return tables, issues, formula_cells, summary


def prepare(path):
    tables, issues, formulas, summary = read_source(path)
    # Prefix fields so pass and case provenance can never silently overwrite one another.
    c = tables['焊接Case'].rename(columns=lambda x: x if x == 'CaseID' else 'Case_' + x)
    p = tables['逐道参数'].rename(columns=lambda x: x if x == 'CaseID' else 'Pass_' + x)
    merged = p.merge(c, on='CaseID', how='left', validate='many_to_one')
    records = []
    for original in merged.to_dict('records'):
        r = {k: (None if pd.isna(v) else v) for k, v in original.items()}
        rr = roles(r['Pass_焊道功能'])
        for role in rr:
            out = dict(r)
            out.update(ObservationID=r['Pass_PassID'], PassID=r['Pass_PassID'], PassRole=role,
                       ContextWeight=1 / len(rr), RoleBasis='Explicit shared setting' if len(rr) > 1 else 'Source role',
                       Thickness=number(r['Case_板厚_mm']), RootGap=number(r['Case_根部间隙_mm']),
                       GrooveAngle=None, Position=position(r['Case_焊接位置_原文']),
                       SourceGroup=r['Case_来源链接'] or r['Case_来源名称'],
                       Approved=r['Case_训练准入'] in ['是', '可用于训练', '可训练'] and r['Pass_训练准入'] in ['是', '可用于训练', '可训练'])
            if '总夹角' in (r['Case_坡口角度定义'] or '') and '未' not in r['Case_坡口角度定义']:
                out['GrooveAngle'] = number(r['Case_坡口角度_deg'])
            nature = r['Pass_参数性质'] or ''
            reason = []
            if r['Case_材料牌号'] != 'AH36' or number(r['Case_ISO4063']) != 136:
                reason.append('Material/process not confirmed AH36/136')
            if r['Case_训练准入'] == '否' or r['Pass_训练准入'] == '否':
                reason.append('Source explicitly excludes training')
            if '候选' in nature:
                reason.append('Engineering starting values, not experiments')
            if any(out[k] is None for k in FEATURES):
                reason.append('Missing or ambiguous input: ' + ','.join(k for k in FEATURES if out[k] is None))
            if nature not in ['实验控制参数', 'PQR实测参数', '文献实测参数', '同行评议实验参数', '研究稿实验参数']:
                reason.append('Not an eligible experimental point')
            if not (r['Pass_核验状态'] or '').startswith('已核验'):
                reason.append('Source not verified in workbook')
            is_range = '范围' in nature or '窗口' in nature
            if '候选' in nature:
                is_range = False
            for target, field in TARGETS.items():
                v = number(r['Pass_' + field])
                out[target + '_SourceValue'] = v
                formula = ('逐道参数', r['Pass_ExcelRow'], field) in formulas
                out[target] = None if formula else v
                out[target + '_Eligible'] = not reason and v is not None and v > 0 and not formula
                out[target + '_Reason'] = '; '.join(reason + (['Formula-based target excluded'] if formula else []) + (['Missing target'] if v is None else []))
            out['DataClass'] = 'TrainableResearch' if any(out[t + '_Eligible'] for t in TARGETS) else ('Constraint' if is_range else 'Auxiliary')
            out['ExclusionReason'] = '; '.join(reason)
            records.append(out)
    training = pd.DataFrame(records)
    summary['contexts'] = len(training)
    summary['class_counts'] = training.groupby('DataClass')['ObservationID'].nunique().to_dict()
    summary['target_coverage'] = {t: {'contexts': int(training[t + '_Eligible'].sum()),
                                    'observations': int(training.loc[training[t + '_Eligible'], 'ObservationID'].nunique()),
                                    'cases': int(training.loc[training[t + '_Eligible'], 'CaseID'].nunique()),
                                    'sources': int(training.loc[training[t + '_Eligible'], 'SourceGroup'].nunique())} for t in TARGETS}
    summary['approved_contexts'] = int(training.Approved.sum())
    summary['cases_without_passes'] = sorted(set(c.CaseID) - set(p.CaseID))
    return tables, training, issues, summary


def save_data(path, destination):
    dest = Path(destination)
    dest.mkdir(parents=True, exist_ok=True)
    tables, training, issues, summary = prepare(path)
    for name, frame in tables.items():
        frame.to_csv(dest / f'{name}.csv', index=False, encoding='utf-8-sig')
    training.to_csv(dest / 'Model_Training_V1.csv', index=False, encoding='utf-8-sig')
    for label in ['TrainableResearch', 'Auxiliary', 'Constraint']:
        training[training.DataClass == label].to_csv(dest / f'{label}.csv', index=False, encoding='utf-8-sig')
    pd.DataFrame(issues).to_csv(dest / 'data_issues.csv', index=False, encoding='utf-8-sig')
    (dest / 'audit.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    return tables, training, issues, summary


def save_v02_audit(path, destination):
    """Create V0.2 eligibility artifacts only. Existing raw data and model files are untouched."""
    dest = Path(destination)
    dest.mkdir(parents=True, exist_ok=True)
    tables, training, issues, summary = prepare(path)
    _, _, formulas, _ = read_source(path)
    v02 = add_v02_eligibility(training, formulas)
    queue = build_promotion_queue(v02)
    audit = v02_audit(v02, queue, summary)
    v02.to_csv(dest / 'Model_Training_V2.csv', index=False, encoding='utf-8-sig')
    queue.to_csv(dest / 'Promotion_Queue_V02.csv', index=False, encoding='utf-8-sig')
    (dest / 'training_eligibility_audit_v02.json').write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    return tables, v02, queue, issues, audit
